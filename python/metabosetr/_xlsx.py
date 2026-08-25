"""Tolerant xlsx loading.

WebIDQ / openxlsx-produced exports can declare worksheet relationships to
drawing parts (``xl/drawings/drawing1.xml``, ``vmlDrawing1.vml``) that are
not actually present in the archive. R's ``readxl`` ignores them, but
``openpyxl`` aborts in ``find_images`` while resolving the dangling target.

The R reader (``read-webidq.R``) is deliberately defensive about vendor
export quirks (padding rows, unvalidated sheet names); this mirrors that
posture for the drawing-relationship quirk so the same files load here.
"""

from __future__ import annotations

import io
import re
import zipfile

import openpyxl

_SHEET_RELS_RE = re.compile(r"^xl/worksheets/_rels/.*\.rels$")
_RELATIONSHIP_RE = re.compile(r"<Relationship\b[^>]*/>")
_TARGET_RE = re.compile(r'Target="([^"]+)"')


def _strip_dangling_rels(rels_xml: str, archive_names: set[str]) -> str:
    """Drop <Relationship> entries whose Target part is missing."""

    def keep(match: "re.Match[str]") -> str:
        block = match.group(0)
        target_match = _TARGET_RE.search(block)
        if target_match is None:
            return block
        # Worksheet rels use paths relative to xl/worksheets/, e.g.
        # "../drawings/drawing1.xml" -> "xl/drawings/drawing1.xml".
        target = target_match.group(1)
        resolved = target.replace("../", "xl/")
        return "" if resolved not in archive_names else block

    return _RELATIONSHIP_RE.sub(keep, rels_xml)


def load_workbook(path: str, read_only: bool = False):
    """Load an xlsx workbook, tolerating dangling drawing relationships.

    Returns an :class:`openpyxl.Workbook`. Falls back to a repackaged
    archive with unresolved worksheet relationships removed when the direct
    load fails on a missing drawing/vmlDrawing target.
    """
    try:
        return openpyxl.load_workbook(path, read_only=read_only)
    except KeyError:
        pass  # dangling relationship target; repair below

    with open(path, "rb") as fh:
        raw = fh.read()
    src = zipfile.ZipFile(io.BytesIO(raw))
    names = set(src.namelist())
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if _SHEET_RELS_RE.match(info.filename):
                data = _strip_dangling_rels(
                    data.decode("utf-8"), names
                ).encode("utf-8")
            dst.writestr(info, data)
    buf.seek(0)
    return openpyxl.load_workbook(buf, read_only=read_only)


def first_sheet_rows(path: str):
    """Return (sheet_name, rows) for the first sheet.

    ``rows`` is a list of row tuples of cell values (``None`` for blanks),
    matching the way the R reader treats the first sheet only (sheet name is
    not validated).
    """
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    rows = [tuple(cell.value for cell in row) for row in ws.iter_rows()]
    return ws.title, rows
